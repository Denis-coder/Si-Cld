# Применяет data/db.json к оригиналу source/original.xlsx -> Perechen_SI.xlsx (оформление и прочие листы сохраняются)
import openpyxl,json,sys,os
from copy import copy
sys.path.insert(0,os.path.dirname(__file__)); from common import *
db=load_db(); out=sys.argv[1] if len(sys.argv)>1 else 'Perechen_SI.xlsx'; wb=openpyxl.load_workbook('source/original.xlsx')
def val(k,v):
    if v=='': return None
    return int(v) if k in('year','interval') and v.isdigit() else v
for name,rows in db.items():
    if name not in wb.sheetnames: continue
    ws=wb[name]; ev=[]; keep={r['x'] for r in rows if r.get('x')}
    for r in rows:
        if r.get('x'):
            for j,k in enumerate(KEYS):
                c=ws.cell(r['x'],j+2)
                if s(c.value)!=r[k]: c.value=val(k,r[k])
    for n in range(2,ws.max_row+1):
        if ws.cell(n,2).value and RX.search(str(ws.cell(n,2).value)) and n not in keep: ev.append((n,0,0,None))
    anchor=1
    for seq,r in enumerate(rows):
        if r.get('x'): anchor=r['x']
        else: ev.append((anchor+1,1,seq,r))
    ev.sort(key=lambda e:(-e[0],e[1],-e[2]))   # снизу вверх; на одной позиции: сначала удаление, потом вставка
    for p,kind,_,r in ev:
        if kind==0: ws.delete_rows(p)
        else:
            ws.insert_rows(p)
            for c in range(1,ws.max_column+1): ws.cell(p,c)._style=copy(ws.cell(p-1,c)._style)
            for j,k in enumerate(KEYS): ws.cell(p,j+2).value=val(k,r[k])
    if ev:
        for n in range(2,ws.max_row+1):
            if ws.cell(n,2).value: ws.cell(n,1).value=f'=IF(B{n}<>"", COUNTA($B$2:B{n}), "")'
wb.save(out); print('ok')
