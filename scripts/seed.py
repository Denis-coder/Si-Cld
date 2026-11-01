# Создаёт data/db.json из source/original.xlsx (запускается вручную: Actions -> seed)
import openpyxl,json,sys,os
sys.path.insert(0,os.path.dirname(__file__)); from common import *
wb=openpyxl.load_workbook('source/original.xlsx'); sheets={}; i=0
for ws in wb:
    if s(ws['H1'].value)!='Место установки СИ' or s(ws['B1'].value)!='Наименование': continue
    rows=[]
    for n,r in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if not r[1] or not RX.search(str(r[1])): continue
        d={k:s(r[j+1]) for j,k in enumerate(KEYS)}; i+=1; d['id']='r%d'%i; d['x']=n; rows.append(d)
    if rows and ws.title!='Замены июль 2026': sheets[ws.title]=rows
save_db(sheets)
print(sum(map(len,sheets.values())),'СИ,',len(sheets),'листов')
